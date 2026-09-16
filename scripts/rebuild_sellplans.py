# Rebuilds SellPlans column D (scenario JSON) with save-day snapshot rows — design.md §C.14.
# Usage: python scripts/rebuild_sellplans.py   (read-only against the Sheet; output goes to backups/)
import datetime as dt
import json
import re
import sys
import time
import urllib.parse
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
TAIPEI = dt.timezone(dt.timedelta(hours=8))
TAX_RATE, FEE_RATE, FEE_MIN = 0.003, 0.001425, 20
# A code is trusted only if its market price agrees with the Sheet's own price for that
# name — src/segments.json has already been caught mapping a holding to another company.
PRICE_TOLERANCE = 0.15
MONTHS = {m: f"{i:02d}" for i, m in enumerate(
    ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], 1)}


# ---------------------------------------------------------------- pure helpers (unit-tested)

def taipei_day(iso):
    """UTC ISO instant -> the Taiwan calendar day it falls on (Taiwan has no DST)."""
    t = dt.datetime.fromisoformat(iso.replace("Z", "+00:00"))
    return t.astimezone(TAIPEI).date().isoformat()


def normalize_date_key(k):
    """Mirror of src/api.js normalizeDateKey: GAS can serve 'Sun Dec 31 2017 08:00:00 GMT+0800 ...'."""
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", k) or k == "closeyest":
        return k
    m = re.match(r"\w{3} (\w{3}) (\d{2}) (\d{4})", k)
    return f"{m.group(3)}-{MONTHS[m.group(1)]}-{m.group(2)}" if m and m.group(1) in MONTHS else k


def net_proceeds(shares, price):
    """Mirror of proceeds() in src/sellplanner.js at a full brokerage rate."""
    if not shares or shares <= 0:
        return 0.0
    gross = shares * price
    return gross - gross * TAX_RATE - max(FEE_MIN, gross * FEE_RATE)


def ledger_sales(trades, lots):
    # Port of realizedSales() in src/sellplanner.js: sold cost = Σ buys − Σ held-lot cost,
    # split across a ticker's sales at the average cost per share sold.
    buy_amount, held_cost, sells = {}, {}, {}
    for lot in lots:
        t = lot["ticker"].strip()
        held_cost[t] = held_cost.get(t, 0) + (lot.get("cost") or 0)
    for tr in trades:
        t = tr["ticker"].strip()
        if tr["type"] == "buy":
            buy_amount[t] = buy_amount.get(t, 0) + (tr.get("amount") or 0)
        elif tr["type"] == "sell":
            sells.setdefault(t, []).append(tr)
    out = []
    for t, ss in sells.items():
        cost = buy_amount.get(t, 0) - held_cost.get(t, 0)
        shares_sold = sum(s.get("shares") or 0 for s in ss)
        per_share = cost / shares_sold if shares_sold > 0 else None
        for s in ss:
            shares = s.get("shares")
            out.append({
                "date": s["date"], "ticker": t, "shares": shares,
                "cost": shares * per_share if shares and per_share is not None else None,
            })
    return out


def position_at_save(ticker, sell_shares, save_day, lots, trades, sales):
    # -> (hold_shares, cost, rolled_back). Today's lots with every trade from the save day on undone.
    # A plan may record sales made before it was saved, so if it sells more than that position,
    # earlier sales are added back newest first until the planned shares are covered.
    hold = sum(l.get("shares") or 0 for l in lots if l["ticker"].strip() == ticker)
    cost = sum(l.get("cost") or 0 for l in lots if l["ticker"].strip() == ticker)
    for s in sales:
        if s["ticker"] == ticker and s["date"] >= save_day and s["shares"] and s["cost"] is not None:
            hold += s["shares"]
            cost += s["cost"]
    for b in trades:
        if b["type"] == "buy" and b["ticker"].strip() == ticker and b["date"] >= save_day:
            hold -= b.get("shares") or 0
            cost -= b.get("amount") or 0
    rolled = 0
    earlier = [s for s in sales if s["ticker"] == ticker and s["date"] < save_day
               and s["shares"] and s["cost"] is not None]
    for s in sorted(earlier, key=lambda s: s["date"], reverse=True):
        if hold >= sell_shares:
            break
        hold += s["shares"]
        cost += s["cost"]
        rolled += s["shares"]
    return hold, cost, rolled


def close_on_or_before(series, day):
    """series: [(yyyy-mm-dd, close)] -> the last (date, close) on or before `day`, else None."""
    best = None
    for d, v in series:
        if d <= day and v is not None and (best is None or d > best[0]):
            best = (d, v)
    return best


def within_tolerance(a, b, tol=PRICE_TOLERANCE):
    return a is not None and b is not None and b > 0 and abs(a - b) / b <= tol


def choose_code(candidates, sheet_price, latest_close):
    # candidates: [(code, source)] by priority. The first whose market price agrees with the Sheet's
    # wins; with no Sheet price to compare, only a code from the user's own dividend records is trusted.
    for code, source in candidates:
        close = latest_close(code)
        if close is None:
            continue
        if sheet_price is None:
            if source == "dividends":
                return code, source
            continue
        if within_tolerance(close, sheet_price):
            return code, source
    return None, None


def as_number(x, places=None):
    if places is not None:
        x = round(x, places)
    return int(x) if float(x).is_integer() else x


def build_row(ticker, sell_shares, price, hold_shares, cost):
    row = {"ticker": ticker, "sellShares": as_number(sell_shares)}
    if price is not None:
        row["sellPrice"] = as_number(price, 2)
    row["holdShares"] = as_number(max(hold_shares, 0))
    row["cost"] = as_number(max(cost, 0), 0)
    return row


def row_figures(row):
    """Net proceeds, realized P/L and % exactly as the planner computes them from a snapshot row."""
    price, n, hold, cost = row.get("sellPrice"), row["sellShares"], row["holdShares"], row["cost"]
    if price is None:
        return None, None, None
    net = net_proceeds(n, price)
    if hold <= 0:
        return net, None, None
    basis = n * cost / hold
    pl = n * price - basis
    return net, pl, (pl / basis * 100 if basis > 0 else None)


# ---------------------------------------------------------------- I/O

def read_config(path):
    text = Path(path).read_text(encoding="utf-8")
    get = lambda name: re.search(rf'{name}\s*=\s*["\']([^"\']+)["\']', text).group(1)
    return get("WEBAPP_URL"), get("API_KEY")


def http_json(url, attempts=3):
    last = None
    for i in range(attempts):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=300) as r:
                return json.loads(r.read().decode("utf-8"))
        except Exception as e:  # noqa: BLE001 — transient Apps Script / network failures are retried
            last = e
            time.sleep(1.5 * (i + 1))
    raise RuntimeError(f"request failed after {attempts} attempts: {last}")


def api_get(base, key, resource, **params):
    q = urllib.parse.urlencode({"key": key, "resource": resource, **params})
    data = http_json(f"{base}?{q}")
    if isinstance(data, dict) and data.get("error"):
        raise RuntimeError(f"Apps Script refused '{resource}': {data['error']}")
    if isinstance(data, list):
        for r in data:
            for f in ("ticker", "name"):
                if isinstance(r.get(f), str):
                    r[f] = r[f].strip()
    return data


def yahoo_series(code, start_day, end_day):
    """Daily closes for a TWSE (.TW) or TPEx (.TWO) code -> [(yyyy-mm-dd, close)] or None."""
    p1 = int(dt.datetime.fromisoformat(start_day).replace(tzinfo=TAIPEI).timestamp())
    p2 = int((dt.datetime.fromisoformat(end_day).replace(tzinfo=TAIPEI) + dt.timedelta(days=1)).timestamp())
    for suffix in (".TW", ".TWO"):
        url = (f"https://query1.finance.yahoo.com/v8/finance/chart/{code}{suffix}"
               f"?period1={p1}&period2={p2}&interval=1d")
        try:
            res = http_json(url, attempts=2)["chart"]["result"][0]
        except Exception:  # noqa: BLE001 — a 404 on .TW simply means try .TWO
            continue
        stamps = res.get("timestamp") or []
        closes = res["indicators"]["quote"][0].get("close") or []
        series = [(dt.datetime.fromtimestamp(ts, TAIPEI).date().isoformat(), c)
                  for ts, c in zip(stamps, closes) if c is not None]
        if series:
            return series
    return None


def main():
    base, key = read_config(ROOT / "config.js")
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    backups = ROOT / "backups"
    backups.mkdir(exist_ok=True)

    print("Reading plans, ledger and prices from your Apps Script (sequentially; this is slow)...")
    plans = api_get(base, key, "sellplans")
    (backups / f"sellplans-original-{stamp}.json").write_text(
        json.dumps(plans, ensure_ascii=False, indent=2), encoding="utf-8")
    lots = api_get(base, key, "heldlots")
    trades = api_get(base, key, "trades")
    dividends = api_get(base, key, "dividends")
    raw_prices = api_get(base, key, "prices", tickers="")

    for tr in trades:
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(tr["date"])):
            sys.exit(f"Unexpected trade date format {tr['date']!r}; refusing to guess.")

    sheet_price = {}
    for name, series in raw_prices.items():
        dated = sorted((normalize_date_key(k), v) for k, v in series.items())
        days = [(d, v) for d, v in dated if re.fullmatch(r"\d{4}-\d{2}-\d{2}", d)]
        sheet_price[name.strip()] = days[-1][1] if days else series.get("closeyest")

    sales = ledger_sales(trades, lots)
    segments = json.loads((ROOT / "src" / "segments.json").read_text(encoding="utf-8"))["tickers"]
    dividend_code = {}
    for d in dividends:
        if d.get("code"):
            dividend_code.setdefault(d["name"], str(d["code"]))

    tickers = sorted({r["ticker"] for p in plans for r in p["rows"]})
    save_days = [taipei_day(p["savedAt"]) for p in plans]
    today = dt.datetime.now(TAIPEI).date().isoformat()
    start = (dt.date.fromisoformat(min(save_days)) - dt.timedelta(days=14)).isoformat() if save_days else today

    series_cache, code_of, problems = {}, {}, []

    def series_for(code):
        if code not in series_cache:
            series_cache[code] = yahoo_series(code, start, today)
        return series_cache[code]

    def latest_close(code):
        s = series_for(code)
        return s[-1][1] if s else None

    print(f"Fetching daily closes for {len(tickers)} stocks from Yahoo Finance...")
    for t in tickers:
        candidates = []
        if dividend_code.get(t):
            candidates.append((dividend_code[t], "dividends"))
        seg = (segments.get(t) or {}).get("code")
        if seg and seg not in [c for c, _ in candidates]:
            candidates.append((str(seg), "segments.json"))
        code, source = choose_code(candidates, sheet_price.get(t), latest_close)
        code_of[t] = code
        if code is None:
            tried = ", ".join(f"{c} ({s})" for c, s in candidates) or "no code on record"
            problems.append(f"{t}: no trustworthy code — tried {tried}; Sheet price {sheet_price.get(t)}")

    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "SellPlans"
    ws.append(["id", "savedAt", "name", "scenario"])
    check = wb.create_sheet("Check")
    check.append(["plan", "saved (Taiwan day)", "ticker", "code", "sellShares", "price date", "sellPrice",
                  "holdShares", "cost", "net proceeds", "realized P/L", "realized P/L %", "note"])
    for cell in (*ws[1], *check[1]):
        cell.font = Font(bold=True)

    print()
    for plan, day in zip(plans, save_days):
        new_rows, tot_net, tot_pl = [], 0.0, 0.0
        for r in plan["rows"]:
            t, n = r["ticker"], r["sellShares"]
            hold, cost, rolled = position_at_save(t, n, day, lots, trades, sales)
            picked = close_on_or_before(series_for(code_of[t]) or [], day) if code_of[t] else None
            row = build_row(t, n, picked[1] if picked else None, hold, cost)
            new_rows.append(row)
            net, pl, pct = row_figures(row)
            tot_net += net or 0
            tot_pl += pl or 0
            notes = []
            if rolled:
                notes.append(f"includes {rolled:g} shares sold before this plan was saved")
            if hold < n:
                notes.append(f"plans {n:g} but only {hold:g} can be accounted for")
            if picked is None:
                notes.append("no save-day price — row keeps no sellPrice")
                problems.append(f"{plan['name']} / {t}: no save-day price")
            check.append([plan["name"], day, t, code_of[t], n, picked[0] if picked else None, row.get("sellPrice"),
                          row["holdShares"], row["cost"], round(net) if net is not None else None,
                          round(pl) if pl is not None else None, round(pct, 1) if pct is not None else None,
                          "; ".join(notes)])
        rebuilt = {k: (new_rows if k == "rows" else v) for k, v in plan.items()}
        ws.append([plan["id"], plan["savedAt"], plan["name"],
                   json.dumps(rebuilt, ensure_ascii=False, separators=(",", ":"))])
        recorded = plan.get("netAtSave")
        check.append([plan["name"], day, "TOTAL", None, None, None, None, None, None, round(tot_net), round(tot_pl),
                      None, f"netAtSave recorded when saved: {round(recorded) if recorded is not None else '—'}"])
        check.append([])
        print(f"  {plan['name']:<24} saved {day}  rows {len(new_rows):>2}  net {tot_net:>12,.0f}  "
              f"P/L {tot_pl:>+12,.0f}  (recorded at save: {recorded:,.0f})" if recorded is not None else
              f"  {plan['name']:<24} saved {day}  rows {len(new_rows):>2}  net {tot_net:>12,.0f}  P/L {tot_pl:>+12,.0f}")

    ws.column_dimensions["D"].width = 120
    out = backups / f"SellPlans-rebuilt-{stamp}.xlsx"
    wb.save(out)
    print(f"\nWrote {out.relative_to(ROOT)}")
    print(f"Original plans backed up to backups/sellplans-original-{stamp}.json")
    if problems:
        print("\nNeeds attention:")
        for p in dict.fromkeys(problems):
            print("  -", p)
    else:
        print("Every row priced.")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
